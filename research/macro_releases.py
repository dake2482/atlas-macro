"""Keyless official BEA and Census release-workbook adapters.

Both agencies publish first-party XLSX workbooks alongside their releases.  We
use those files when free API credentials are unavailable, retain workbook
hashes, and keep estimate/revision labels in observation metadata.  Nothing in
this module treats retrieval time as the economic-data vintage.
"""

from __future__ import annotations

import hashlib
import re
import zipfile
from datetime import datetime
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from io import BytesIO
from typing import Any
from urllib.parse import urljoin, urlparse

from openpyxl import load_workbook

from .providers import HTTPProvider, ProviderResult

BEA_GDP_PAGE = "https://www.bea.gov/data/gdp/gross-domestic-product"
BEA_VINTAGE_WORKBOOK = (
    "https://apps.bea.gov/national/xls/gdp-gdi-vintage-history.xlsx"
)
CENSUS_MARTS_INDEX = "https://www2.census.gov/retail/releases/historical/marts/"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class _LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.links: list[tuple[str, str]] = []
        self._href: str | None = None
        self._text: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() == "a":
            self._href = dict(attrs).get("href")
            self._text = []

    def handle_data(self, data: str) -> None:
        if self._href is not None:
            self._text.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() == "a" and self._href is not None:
            self.links.append((self._href, " ".join("".join(self._text).split())))
            self._href = None
            self._text = []


def _decimal(value: Any) -> Decimal | None:
    if value is None:
        return None
    text = str(value).strip().replace(",", "")
    if not text or set(text) <= {"."} or text.upper() in {"NA", "N/A", "(NA)", "(*)"}:
        return None
    try:
        parsed = Decimal(text)
    except (InvalidOperation, ValueError):
        return None
    return parsed if parsed.is_finite() else None


def _quarter_start(period: str) -> str | None:
    match = re.fullmatch(r"(\d{4})Q([1-4])", str(period).strip().upper())
    if not match:
        return None
    year, quarter = (int(value) for value in match.groups())
    return f"{year:04d}-{(quarter - 1) * 3 + 1:02d}-01"


def _previous_quarter(period: str) -> str:
    year, quarter = (int(value) for value in re.fullmatch(r"(\d{4})Q([1-4])", period).groups())
    if quarter == 1:
        return f"{year - 1}Q4"
    return f"{year}Q{quarter - 1}"


def _release_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value or "").strip()
    match = re.search(r"[A-Z][a-z]{2,8}\s+\d{1,2},\s+\d{4}", text)
    if not match:
        return None
    for date_format in ("%b %d, %Y", "%B %d, %Y"):
        try:
            return datetime.strptime(match.group(0), date_format).date().isoformat()
        except ValueError:
            continue
    return None


def _estimate_round(value: Any) -> str | None:
    match = re.search(
        r"\((Advance|Second|Third|Revised)\s+Estimate\)",
        str(value or ""),
        re.IGNORECASE,
    )
    return match.group(1).title() if match else None


def _xlsx_bytes(content: bytes, *, max_expanded_bytes: int) -> bytes:
    if not content.startswith(b"PK"):
        raise ValueError("response is not an XLSX ZIP archive")
    with zipfile.ZipFile(BytesIO(content)) as archive:
        total = sum(item.file_size for item in archive.infolist())
        if total > max_expanded_bytes:
            raise ValueError("XLSX exceeded configured expanded-size limit")
        if "[Content_Types].xml" not in archive.namelist():
            raise ValueError("XLSX content-types manifest is missing")
    return content


def _artifact(url: str, content: bytes, content_type: str) -> dict[str, Any]:
    return {
        "url": url,
        "sha256": hashlib.sha256(content).hexdigest(),
        "size": len(content),
        "content_type": content_type,
    }


class _ReleaseWorkbookProvider(HTTPProvider):
    max_html_bytes = 5 * 1024 * 1024
    max_workbook_bytes = 4 * 1024 * 1024
    max_expanded_bytes = 48 * 1024 * 1024
    allowed_hosts: frozenset[str] = frozenset()

    def _download(self, url: str, *, expected: str) -> tuple[bytes, str, str]:
        limit = self.max_html_bytes if expected == "html" else self.max_workbook_bytes
        chunks: list[bytes] = []
        size = 0
        with self.client.stream("GET", url) as response:
            response.raise_for_status()
            final_url = urlparse(str(response.url))
            if (
                final_url.scheme != "https"
                or not final_url.hostname
                or final_url.hostname.lower() not in self.allowed_hosts
            ):
                raise ValueError("download redirected outside the official host allowlist")
            for chunk in response.iter_bytes():
                size += len(chunk)
                if size > limit:
                    raise ValueError(f"{expected} response exceeded configured size limit")
                chunks.append(chunk)
            content_type = response.headers.get("content-type", "").split(";", 1)[0].lower()
            last_modified = response.headers.get("last-modified", "")
        content = b"".join(chunks)
        if expected == "html":
            if content_type and "html" not in content_type:
                raise ValueError(f"unexpected HTML content type: {content_type}")
        else:
            _xlsx_bytes(content, max_expanded_bytes=self.max_expanded_bytes)
            if content_type and content_type not in {
                XLSX_CONTENT_TYPE,
                "application/octet-stream",
                "application/zip",
            }:
                raise ValueError(f"unexpected XLSX content type: {content_type}")
        return (
            content,
            content_type or ("text/html" if expected == "html" else XLSX_CONTENT_TYPE),
            last_modified,
        )


class BEAGDPReleaseProvider(_ReleaseWorkbookProvider):
    """Parse BEA's current GDP release and official vintage-history workbook."""

    key = "bea-release"
    base_url = "https://www.bea.gov"
    allowed_hosts = frozenset({"bea.gov", "www.bea.gov", "apps.bea.gov"})

    GROWTH_SERIES = {
        "Personal consumption expenditures": "BEA-DPCERL",
        "Gross private domestic investment": "BEA-GPDI-GROWTH",
        "Fixed investment": "BEA-FIXED-INVESTMENT-GROWTH",
        "Exports": "BEA-EXPORTS-GROWTH",
        "Imports": "BEA-IMPORTS-GROWTH",
        "Government consumption expenditures and gross investment": "BEA-GOVERNMENT-GROWTH",
    }
    CONTRIBUTION_SERIES = {
        "Personal consumption expenditures": "BEA-PCE-CONTRIBUTION",
        "Gross private domestic investment": "BEA-GPDI-CONTRIBUTION",
        "Net exports of goods and services": "BEA-NET-EXPORTS-CONTRIBUTION",
        "Government consumption expenditures and gross investment": "BEA-GOVERNMENT-CONTRIBUTION",
    }

    def gdp_pce(self) -> ProviderResult:
        dataset = "gdp-release-workbooks"
        try:
            page, page_type, _ = self._download(BEA_GDP_PAGE, expected="html")
            comparison_url = self._comparison_url(page)
            vintage, vintage_type, vintage_modified = self._download(
                BEA_VINTAGE_WORKBOOK, expected="xlsx"
            )
            comparison, comparison_type, comparison_modified = self._download(
                comparison_url, expected="xlsx"
            )
            records, release_metadata = self._parse_vintage(vintage)
            component_records, component_metadata = self._parse_components(comparison)
            records.extend(component_records)
        except Exception as exc:
            return ProviderResult.failure(self.key, dataset, f"{type(exc).__name__}: {exc}")

        artifacts = [
            _artifact(BEA_GDP_PAGE, page, page_type),
            _artifact(BEA_VINTAGE_WORKBOOK, vintage, vintage_type),
            _artifact(comparison_url, comparison, comparison_type),
        ]
        return ProviderResult(
            provider=self.key,
            dataset=dataset,
            records=records,
            metadata={
                "source_url": BEA_GDP_PAGE,
                "vintage_workbook_url": BEA_VINTAGE_WORKBOOK,
                "comparison_workbook_url": comparison_url,
                "vintage_last_modified": vintage_modified,
                "comparison_last_modified": comparison_modified,
                "artifacts": artifacts,
                "vintage_policy": "latest published vintage for each estimate quarter",
                "unit_policy": "source workbook values retained",
                "attribution": "U.S. Bureau of Economic Analysis",
                **release_metadata,
                **component_metadata,
            },
        )

    @staticmethod
    def _comparison_url(page: bytes) -> str:
        parser = _LinkParser()
        parser.feed(page.decode("utf-8", errors="replace"))
        candidates = []
        for href, label in parser.links:
            absolute = urljoin(BEA_GDP_PAGE, href)
            parsed = urlparse(absolute)
            if (
                parsed.scheme == "https"
                and parsed.hostname in {"bea.gov", "www.bea.gov"}
                and parsed.path.lower().endswith(".xlsx")
                and ("Historical Comparisons" in label or re.search(r"/hist[^/]+\.xlsx$", parsed.path, re.I))
            ):
                candidates.append(absolute)
        if len(set(candidates)) != 1:
            raise ValueError("BEA page must expose exactly one historical-comparisons workbook")
        return candidates[0]

    @staticmethod
    def _parse_vintage(content: bytes) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        if "Vintage History" not in workbook.sheetnames:
            raise ValueError("BEA vintage workbook is missing the expected sheet")
        sheet = workbook["Vintage History"]
        updated = ""
        records: list[dict[str, Any]] = []
        current_period: str | None = None
        for row in sheet.iter_rows(values_only=True):
            values = list(row) + [None] * 7
            if not updated and isinstance(values[0], str) and values[0].startswith("Last Updated"):
                updated = values[0].removeprefix("Last Updated").strip()
            if isinstance(values[0], str) and _quarter_start(values[0]):
                current_period = values[0].strip().upper()
                continue
            if current_period is None:
                continue
            nominal_gdp = _decimal(values[2])
            real_gdp = _decimal(values[4])
            if not values[1] or nominal_gdp is None or real_gdp is None or not values[6]:
                continue
            value_date = _quarter_start(current_period)
            metadata = {
                "estimate_quarter": current_period,
                "vintage_label": str(values[1]).strip(),
                "estimate_round": str(values[1]).strip(),
                "source_revision_date": _release_date(values[6]),
                "source_revision_text": str(values[6]).strip(),
                "workbook_last_updated": updated,
                "seasonal_adjustment": "SAAR",
            }
            records.extend(
                [
                    {"series_id": "BEA-A191RL", "date": value_date, "value": real_gdp, "metadata": {**metadata, "unit": "percent change from preceding period"}},
                    {"series_id": "BEA-GDP-NOMINAL-SAAR", "date": value_date, "value": nominal_gdp, "metadata": {**metadata, "unit": "USD billions, current dollars"}},
                ]
            )
            nominal_gdi = _decimal(values[3])
            real_gdi = _decimal(values[5])
            if nominal_gdi is not None:
                records.append({"series_id": "BEA-GDI-NOMINAL-SAAR", "date": value_date, "value": nominal_gdi, "metadata": {**metadata, "unit": "USD billions, current dollars"}})
            if real_gdi is not None:
                records.append({"series_id": "BEA-GDI-REAL-GROWTH-SAAR", "date": value_date, "value": real_gdi, "metadata": {**metadata, "unit": "percent change from preceding period"}})
            current_period = None
        if not records:
            raise ValueError("BEA vintage workbook yielded no GDP observations")
        return records, {"workbook_last_updated": updated, "quarter_count": len({item["date"] for item in records})}

    def _parse_components(
        self, content: bytes
    ) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        if "GDPhistQ" not in workbook.sheetnames:
            raise ValueError("BEA comparisons workbook is missing GDPhistQ")
        sheet = workbook["GDPhistQ"]
        rows = [list(row) + [None] * 8 for row in sheet.iter_rows(values_only=True)]
        growth_index = next(
            (
                index
                for index, values in enumerate(rows)
                if "Comparisons -- Percent Change from Preceding Period"
                in str(values[0] or "").replace("\n", " ")
            ),
            None,
        )
        contribution_index = next(
            (
                index
                for index, values in enumerate(rows)
                if "Comparisons -- Contributions to Percent Change"
                in str(values[0] or "").replace("\n", " ")
            ),
            None,
        )
        if growth_index is None or contribution_index is None:
            raise ValueError("BEA comparisons workbook is missing growth or contribution sections")
        growth_title = str(rows[growth_index][0] or "")
        match = re.search(r"(\d{4}Q[1-4])", growth_title)
        if not match:
            raise ValueError("BEA comparisons workbook has no estimate quarter")
        period = match.group(1)
        previous_period = _previous_quarter(period)
        release_date = _release_date(sheet.cell(1, 1).value)
        estimate_round = _estimate_round(growth_title)

        def parse_section(
            section_rows: list[list[Any]],
            series: dict[str, str],
            *,
            unit: str,
            section_name: str,
        ) -> tuple[list[dict[str, Any]], set[str]]:
            records: list[dict[str, Any]] = []
            found: set[str] = set()
            context = ""
            for values in section_rows:
                label = str(values[0] or "").strip()
                if label in {
                    "Personal consumption expenditures",
                    "Gross private domestic investment",
                    "Exports",
                    "Imports",
                    "Government consumption expenditures and gross investment",
                    "Net exports of goods and services",
                }:
                    context = label
                series_id = series.get(label)
                if section_name == "growth" and context == "Personal consumption expenditures":
                    if label == "Goods":
                        series_id = "BEA-PCE-GOODS-GROWTH"
                    elif label == "Services":
                        series_id = "BEA-PCE-SERVICES-GROWTH"
                current_value = _decimal(values[1])
                if not series_id or current_value is None:
                    continue
                metadata = {
                    "estimate_quarter": period,
                    "estimate_round": estimate_round,
                    "source_revision_date": release_date,
                    "seasonal_adjustment": "SAAR",
                    "unit": unit,
                    "component_label": label,
                    "comparison_section": section_name,
                }
                records.append(
                    {
                        "series_id": series_id,
                        "date": _quarter_start(period),
                        "value": current_value,
                        "metadata": metadata,
                    }
                )
                for period_index, value_index in ((2, 3), (4, 5), (6, 7)):
                    if str(values[period_index] or "").strip().upper() == previous_period:
                        previous_value = _decimal(values[value_index])
                        if previous_value is not None:
                            records.append(
                                {
                                    "series_id": series_id,
                                    "date": _quarter_start(previous_period),
                                    "value": previous_value,
                                    "metadata": {
                                        **metadata,
                                        "estimate_quarter": previous_period,
                                    },
                                }
                            )
                        break
                found.add(label)
            return records, found

        growth_records, growth_found = parse_section(
            rows[growth_index + 1 : contribution_index],
            self.GROWTH_SERIES,
            unit="percent change from preceding period",
            section_name="growth",
        )
        contribution_end = next(
            (
                index
                for index in range(contribution_index + 1, len(rows))
                if "Comparisons --" in str(rows[index][0] or "").replace("\n", " ")
            ),
            len(rows),
        )
        contribution_records, contribution_found = parse_section(
            rows[contribution_index + 1 : contribution_end],
            self.CONTRIBUTION_SERIES,
            unit="percentage points contribution to real GDP growth",
            section_name="contribution",
        )
        required_growth = {"Personal consumption expenditures", "Goods", "Services"}
        required_contributions = set(self.CONTRIBUTION_SERIES)
        if not required_growth <= growth_found:
            raise ValueError("BEA comparisons workbook is missing required PCE growth components")
        if not required_contributions <= contribution_found:
            raise ValueError("BEA comparisons workbook is missing required GDP contributions")
        return [*growth_records, *contribution_records], {
            "comparison_quarter": period,
            "comparison_release_date": release_date,
            "comparison_estimate_round": estimate_round,
        }


class CensusMARTSReleaseProvider(_ReleaseWorkbookProvider):
    """Parse Census Monthly Retail Trade official release workbooks without an API key."""

    key = "census-release"
    base_url = "https://www2.census.gov"
    allowed_hosts = frozenset({"www2.census.gov"})

    def monthly_retail_sales(self) -> ProviderResult:
        dataset = "marts:retail-food-services"
        try:
            index, index_type, _ = self._download(CENSUS_MARTS_INDEX, expected="html")
            workbook_url = self._latest_workbook_url(index)
            content, content_type, last_modified = self._download(workbook_url, expected="xlsx")
            records, metadata = self._parse_workbook(content)
        except Exception as exc:
            return ProviderResult.failure(self.key, dataset, f"{type(exc).__name__}: {exc}")
        artifacts = [
            _artifact(CENSUS_MARTS_INDEX, index, index_type),
            _artifact(workbook_url, content, content_type),
        ]
        return ProviderResult(
            provider=self.key,
            dataset=dataset,
            records=records,
            metadata={
                "source_url": CENSUS_MARTS_INDEX,
                "workbook_url": workbook_url,
                "workbook_last_modified": last_modified,
                "artifacts": artifacts,
                "unit": "USD millions",
                "vintage_policy": "latest official advance/preliminary/revised release workbook",
                "attribution": "U.S. Census Bureau",
                **metadata,
            },
        )

    @staticmethod
    def _latest_workbook_url(index: bytes) -> str:
        parser = _LinkParser()
        parser.feed(index.decode("latin-1", errors="replace"))
        candidates: list[tuple[int, int, str]] = []
        for href, _ in parser.links:
            match = re.fullmatch(r"rs(\d{2})(\d{2})\.xlsx", href.strip(), re.I)
            if not match:
                continue
            short_year, month = (int(value) for value in match.groups())
            year = 2000 + short_year if short_year < 80 else 1900 + short_year
            if 1 <= month <= 12:
                candidates.append((year, month, href))
        if not candidates:
            raise ValueError("Census MARTS directory exposed no release workbook")
        _, _, href = max(candidates)
        return urljoin(CENSUS_MARTS_INDEX, href)

    @staticmethod
    def _month_date(year: int, label: Any) -> str | None:
        text = re.sub(r"[^A-Za-z]", "", str(label or ""))[:3].title()
        try:
            month = datetime.strptime(text, "%b").month
        except ValueError:
            return None
        return f"{year:04d}-{month:02d}-01"

    @staticmethod
    def _month_from_text(value: Any) -> str | None:
        match = re.search(r"([A-Z][a-z]{2})\.?\s+(\d{4})", str(value or ""))
        if not match:
            return None
        try:
            month = datetime.strptime(match.group(1), "%b").month
        except ValueError:
            return None
        return f"{int(match.group(2)):04d}-{month:02d}-01"

    @classmethod
    def _parse_workbook(cls, content: bytes) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        if not {"Table 1.", "Table 2."} <= set(workbook.sheetnames):
            raise ValueError("Census MARTS workbook is missing required tables")
        sales_sheet = workbook["Table 1."]
        rows = [list(row) for row in sales_sheet.iter_rows(values_only=True)]
        adjusted_column = None
        for row in rows[:15]:
            for index, value in enumerate(row):
                if isinstance(value, str) and value.strip().lower().startswith("adjusted"):
                    adjusted_column = index
                    break
            if adjusted_column is not None:
                break
        if adjusted_column is None:
            raise ValueError("Census MARTS adjusted-sales columns were not found")

        year_by_column: dict[int, int] = {}
        month_by_column: dict[int, Any] = {}
        status_by_column: dict[int, str] = {}
        active_year: int | None = None
        for row in rows[:15]:
            for index in range(adjusted_column, len(row)):
                value = row[index]
                if isinstance(value, int) and 1900 <= value <= 2100:
                    active_year = value
                if active_year is not None:
                    year_by_column.setdefault(index, active_year)
                if isinstance(value, str) and re.search(r"[A-Za-z]{3}", value):
                    if cls._month_date(active_year or 0, value):
                        month_by_column[index] = value
                if isinstance(value, str) and re.fullmatch(r"\([apr]\)", value.strip(), re.I):
                    status_by_column[index] = value.strip().lower()

        total_row = None
        previous_label = ""
        for row in rows:
            label = str(row[1] or "") if len(row) > 1 else ""
            if "Retail & food services" in previous_label and label.strip().lower().startswith("total"):
                total_row = row
                break
            if label.strip():
                previous_label = label
        if total_row is None:
            raise ValueError("Census MARTS total sales row was not found")

        records = []
        dates: list[str] = []
        status_by_date: dict[str, str] = {}
        for index in range(adjusted_column, min(adjusted_column + 3, len(total_row))):
            value = _decimal(total_row[index])
            year = year_by_column.get(index)
            value_date = cls._month_date(year or 0, month_by_column.get(index))
            if value is None or value_date is None:
                continue
            dates.append(value_date)
            status_by_date[value_date] = status_by_column.get(index, "")
            records.append(
                {
                    "series_id": "CENSUS-MRTS-44X72-SM-SA",
                    "date": value_date,
                    "value": value,
                    "metadata": {
                        "unit": "USD millions",
                        "seasonally_adjusted": True,
                        "estimate_status": status_by_column.get(index, ""),
                        "category": "Retail and food services, total",
                    },
                }
            )
        if len(records) < 2:
            raise ValueError("Census MARTS workbook yielded insufficient monthly sales")
        if dates != sorted(set(dates), reverse=True):
            raise ValueError("Census MARTS adjusted-sales months are not unique and descending")

        change_sheet = workbook["Table 2."]
        change_rows = [list(row) for row in change_sheet.iter_rows(values_only=True)]
        change_total = None
        previous_label = ""
        for row in change_rows:
            label = str(row[1] or "") if len(row) > 1 else ""
            if "Retail & food services" in previous_label and label.strip().lower().startswith("total"):
                change_total = row
                break
            if label.strip():
                previous_label = label
        if change_total is None:
            raise ValueError("Census MARTS change row was not found")
        change_columns: dict[str, int] = {}
        for row in change_rows[:15]:
            for index, value in enumerate(row):
                value_date = cls._month_from_text(value)
                if value_date and re.search(
                    r"\b(Advance|Preliminary|Revised)\b", str(value), re.IGNORECASE
                ):
                    change_columns.setdefault(value_date, index)
        if not all(value_date in change_columns for value_date in dates[:2]):
            raise ValueError("Census MARTS change-table month headers do not match sales table")

        for position, value_date in enumerate(dates[:2]):
            mom_index = change_columns[value_date]
            yoy_index = mom_index + 1
            comparison_row = next(
                (
                    row
                    for row in change_rows[:15]
                    if mom_index < len(row)
                    and cls._month_from_text(row[mom_index]) not in {None, value_date}
                ),
                None,
            )
            if comparison_row is None or yoy_index >= len(comparison_row):
                raise ValueError("Census MARTS change-table comparison headers are incomplete")
            prior_month = cls._month_from_text(comparison_row[mom_index])
            year_ago = cls._month_from_text(comparison_row[yoy_index])
            expected_prior = dates[position + 1]
            expected_year_ago = f"{int(value_date[:4]) - 1:04d}{value_date[4:]}"
            if prior_month != expected_prior:
                raise ValueError("Census MARTS current month-over-month header is inconsistent")
            if year_ago != expected_year_ago:
                raise ValueError("Census MARTS year-over-year header is inconsistent")
            for series_id, value_index in (
                ("CENSUS-MRTS-44X72-SM-SA-MOM", mom_index),
                ("CENSUS-MRTS-44X72-SM-SA-YOY", yoy_index),
            ):
                value = _decimal(change_total[value_index])
                if value is not None:
                    records.append(
                        {
                            "series_id": series_id,
                            "date": value_date,
                            "value": value,
                            "metadata": {
                                "unit": "percent",
                                "seasonally_adjusted": True,
                                "estimate_status": status_by_date.get(value_date, ""),
                                "category": "Retail and food services, total",
                            },
                        }
                    )
        return records, {
            "latest_value_date": dates[0],
            "months_in_workbook": dates,
            "precision_policy": "retain workbook values without additional rounding",
        }
